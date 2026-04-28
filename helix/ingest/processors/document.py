# Copyright (c) 2026 HELIX. All rights reserved.
# HELIX Personal Intelligence OS
"""Document processor for Office files."""

from pathlib import Path
from helix.ingest.processors.base import BaseProcessor, RawContent
import logging
import asyncio

logger = logging.getLogger(__name__)

class DocumentProcessor(BaseProcessor, mime_patterns=[
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "application/msword"
]):
    """Extracts text from Office documents."""

    def can_handle(self, mime_type: str, extension: str) -> bool:
        return mime_type.startswith("application/vnd") or mime_type == "application/msword"

    async def extract(self, path: Path) -> RawContent:
        """Extract raw content from a document."""
        logger.info(f"Extracting Document: {path}")
        text_content = ""
        ext = path.suffix.lower()

        try:
            if ext == ".docx":
                import docx
                def _read_docx() -> str:
                    doc = docx.Document(str(path))
                    return "\n".join([p.text for p in doc.paragraphs])
                text_content = await asyncio.to_thread(_read_docx)

            elif ext == ".xlsx":
                import openpyxl
                def _read_xlsx() -> str:
                    wb = openpyxl.load_workbook(path, data_only=True)
                    rows = []
                    # Just read the first sheet to avoid massive text overflow
                    if wb.sheetnames:
                        ws = wb[wb.sheetnames[0]]
                        for row in ws.iter_rows(values_only=True, max_row=100):
                            rows.append("\t".join([str(cell) if cell is not None else "" for cell in row]))
                    return "\n".join(rows)
                text_content = await asyncio.to_thread(_read_xlsx)

            elif ext == ".pptx":
                import pptx
                def _read_pptx() -> str:
                    prs = pptx.Presentation(str(path))
                    text_runs = []
                    for slide in prs.slides:
                        for shape in slide.shapes:
                            if hasattr(shape, "text"):
                                text_runs.append(shape.text)
                    return "\n".join(text_runs)
                text_content = await asyncio.to_thread(_read_pptx)
            else:
                text_content = f"Unsupported or legacy document format for {path.name}"

        except Exception as e:
            logger.error(f"Failed to process document {path}: {e}")
            text_content = f"Error processing document: {e}"

        truncated_text = text_content[:8000]

        return RawContent(
            text=truncated_text,
            full_text=text_content,
            frames=[],
            audio_path=None,
            metadata={"source": "DocumentProcessor"},
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            file_path=path
        )
